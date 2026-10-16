import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использовать openpyxl ИЛИ xlsxwriter.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    csv_path = Path(csv_path)
    xlsx_path = Path(xlsx_path)

    if not csv_path.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    if csv_path.suffix.lower() != ".csv":
        raise ValueError("Входной файл должен иметь расширение .csv")

    if xlsx_path.suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен иметь расширение .xlsx")

    with csv_path.open(encoding="utf-8", newline="") as file:
        try:
            reader = csv.reader(file, strict=True)
            header = next(reader, None)

            if not header:
                raise ValueError("Пустой CSV или отсутствует заголовок")

            if any(not name.strip() for name in header):
                raise ValueError("Названия столбцов не должны быть пустыми")

            if len(header) != len(set(header)):
                raise ValueError("Названия столбцов не должны повторяться")

            rows = [header]

            for row in reader:
                if not row:
                    continue

                if len(row) != len(header):
                    raise ValueError("Количество полей не совпадает с заголовком")

                rows.append(row)

        except csv.Error as error:
            raise ValueError("Некорректный формат CSV") from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            cell = sheet.cell(
                row=row_number,
                column=column_number,
                value=value)
            cell.data_type = "s"

    for column_number in range(1, len(header) + 1):
        max_length = 0

        for row in rows:
            value = row[column_number - 1]

            for line in value.splitlines():
                max_length = max(max_length, len(line))

        column_letter = get_column_letter(column_number)
        sheet.column_dimensions[column_letter].width = min(
            255, max(8, max_length + 2))
    workbook.save(xlsx_path)
    workbook.close()